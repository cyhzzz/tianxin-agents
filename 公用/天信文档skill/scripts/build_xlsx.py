#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""天信 Excel 表样生成器

用途：把底稿类表格（勾稽校验、底稿索引、抽样、税负测算、工资）生成统一样式的 .xlsx，
避免每次手工拼表导致列宽、字体、边框、打印设置各人不一。

用法：
    python scripts/build_xlsx.py --template 勾稽校验表 --output 勾稽校验表.xlsx
    python scripts/build_xlsx.py --spec my.json --output 自定义.xlsx
    python scripts/build_xlsx.py --verify 勾稽校验表.xlsx

依赖：openpyxl（pip install openpyxl）。没有 openpyxl 时脚本直接报错退出，不会产出半成品。

spec.json 结构：
{
  "title": "表名（可选，占一行合并单元格）",
  "note": "表下说明（可选）",
  "sheets": [
    {"name": "表1",
     "columns": [{"header": "编号", "width": 8}, {"header": "金额", "width": 14, "fmt": "#,##0.00"}],
     "rows": [["A1", 1234.5]],
     "freeze": true,
     "total": false}
  ]
}
"""

import argparse
import json
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("缺少 openpyxl，先装：pip install openpyxl")

NAVY = "1D3A5E"      # 文档轨品牌深蓝（与 assets/report.css 的 --tx-navy 一致）
BORDER = "E0E0E0"
ZEBRA = "F4F4F4"
MUTED = "525252"
CN_FONT = "宋体"

THIN = Side(style="thin", color=BORDER)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header_style(cell):
    cell.font = Font(name=CN_FONT, size=10.5, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BOX


def _body_style(cell, zebra=False, center=False):
    cell.font = Font(name=CN_FONT, size=10.5)
    cell.alignment = Alignment(
        horizontal="center" if center else "left", vertical="center", wrap_text=True
    )
    cell.border = BOX
    if zebra:
        cell.fill = PatternFill("solid", fgColor=ZEBRA)


def build_sheet(wb, spec, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = spec.get("name", "表1")[:31]

    cols = spec.get("columns") or []
    if not cols:
        raise ValueError(f"工作表「{ws.title}」没有 columns")
    ncol = len(cols)

    row = 1
    if spec.get("title"):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        c = ws.cell(row=1, column=1, value=spec["title"])
        c.font = Font(name=CN_FONT, size=14, bold=True, color=NAVY)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 26
        row = 2

    header_row = row
    for i, col in enumerate(cols, start=1):
        cell = ws.cell(row=header_row, column=i, value=col.get("header", ""))
        _header_style(cell)
        ws.column_dimensions[get_column_letter(i)].width = col.get("width", 14)
    ws.row_dimensions[header_row].height = 28

    center_headers = {"编号", "序号", "结论", "金额", "合计", "数量", "比例", "天数", "页码"}
    for r, data in enumerate(spec.get("rows") or [], start=header_row + 1):
        zebra = (r - header_row) % 2 == 0
        for i in range(1, ncol + 1):
            value = data[i - 1] if i - 1 < len(data) else None
            cell = ws.cell(row=r, column=i, value=value)
            head = cols[i - 1].get("header", "")
            _body_style(cell, zebra=zebra, center=head in center_headers)
            fmt = cols[i - 1].get("fmt")
            if fmt:
                cell.number_format = fmt
        ws.row_dimensions[r].height = 22

    if spec.get("total"):
        tr = header_row + len(spec.get("rows") or []) + 1
        for i, col in enumerate(cols, start=1):
            cell = ws.cell(row=tr, column=i)
            _body_style(cell, center=True)
            cell.font = Font(name=CN_FONT, size=10.5, bold=True)
            if i == 1:
                cell.value = "合计"
            elif col.get("sum"):
                letter = get_column_letter(i)
                cell.value = f"=SUM({letter}{header_row + 1}:{letter}{tr - 1})"
                cell.number_format = col.get("fmt", "#,##0.00")

    if spec.get("note"):
        nr = header_row + len(spec.get("rows") or []) + (2 if spec.get("total") else 1)
        ws.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=ncol)
        c = ws.cell(row=nr, column=1, value=spec["note"])
        c.font = Font(name=CN_FONT, size=9, color=MUTED)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    if spec.get("freeze", True):
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # 打印：A4、按宽度缩成一页、重复表头、页脚页码
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = spec.get("orientation", "landscape")
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    ws.oddFooter.center.size = 9
    ws.oddFooter.center.font = CN_FONT
    return ws


# ── 内置表样 ───────────────────────────────────────────────

TEMPLATES = {
    "勾稽校验表": {
        "title": "四表一注勾稽校验表",
        "note": "逐条校验，不平写差异原因；差额为 0 也要留痕。核对依据见 references/11-Excel表样与公式.md。",
        "sheets": [
            {
                "name": "勾稽校验",
                "columns": [
                    {"header": "编号", "width": 8},
                    {"header": "校验项目", "width": 34},
                    {"header": "左值", "width": 16, "fmt": "#,##0.00"},
                    {"header": "右值", "width": 16, "fmt": "#,##0.00"},
                    {"header": "差额", "width": 14, "fmt": "#,##0.00"},
                    {"header": "结论", "width": 10},
                    {"header": "差异原因 / 处理", "width": 32},
                    {"header": "复核人", "width": 10},
                ],
                "rows": [
                    ["A1", "资产总计 = 负债合计 + 所有者权益合计", None, None, None, "", "", ""],
                    ["A2", "流动资产 + 非流动资产 = 资产总计", None, None, None, "", "", ""],
                    ["A3", "营业利润逐级计算", None, None, None, "", "", ""],
                    ["A4", "利润总额 = 营业利润 + 营业外收入 − 营业外支出", None, None, None, "", "", ""],
                    ["A5", "净利润 = 利润总额 − 所得税费用", None, None, None, "", "", ""],
                    ["A6", "经营 + 投资 + 筹资 + 汇率影响 = 现金净增加额", None, None, None, "", "", ""],
                    ["A7", "期初现金 + 净增加额 = 期末现金", None, None, None, "", "", ""],
                    ["B1", "未分配利润期末 = 期初 + 净利润 − 提取盈余公积 − 分配股利", None, None, None, "", "", ""],
                    ["B2", "盈余公积期末 = 期初 + 本期提取 − 弥补亏损/转增", None, None, None, "", "", ""],
                    ["B3", "综合收益总额 = 净利润 + 其他综合收益税后净额", None, None, None, "", "", ""],
                    ["B4", "权益变动表年末余额 = 资产负债表所有者权益", None, None, None, "", "", ""],
                    ["B5", "期末现金 = 货币资金 − 受限资金 − 三个月以上定期", None, None, None, "", "", ""],
                    ["B6", "间接法起点 = 利润表净利润", None, None, None, "", "", ""],
                    ["B7", "间接法经营净额 = 直接法经营净额", None, None, None, "", "", ""],
                    ["B8", "所得税费用 = 当期所得税 + 递延所得税（不含权益部分）", None, None, None, "", "", ""],
                    ["B9", "净利润 = 归母净利润 + 少数股东损益", None, None, None, "", "", ""],
                    ["C1", "附注项目合计 = 报表项目金额", None, None, None, "", "", ""],
                    ["C2", "附注账龄各段之和 = 项目余额", None, None, None, "", "", ""],
                    ["C3", "坏账准备：期初 + 计提 − 收回核销 = 期末", None, None, None, "", "", ""],
                    ["C4", "固定资产：原值期初 + 增加 − 减少 = 期末", None, None, None, "", "", ""],
                    ["C5", "应交税费：期初 + 应交 − 已交 = 期末", None, None, None, "", "", ""],
                    ["C6", "现金流量表补充资料加回数 = 附注对应项目", None, None, None, "", "", ""],
                    ["C7", "附注收入分解合计 = 利润表营业收入", None, None, None, "", "", ""],
                    ["C8", "前五名合计 ≤ 项目总额", None, None, None, "", "", ""],
                    ["C9", "关联方交易金额 ≤ 收入/成本", None, None, None, "", "", ""],
                    ["C10", "分部收入合计 = 营业收入（差异需解释）", None, None, None, "", "", ""],
                ],
            }
        ],
    },
    "底稿索引表": {
        "title": "审计底稿索引表",
        "note": "索引号规则 ZT-{年份}-{客户简称}-{类}-{循环}-{序号}；无索引 = 没做。",
        "sheets": [
            {
                "name": "索引",
                "columns": [
                    {"header": "索引号", "width": 26},
                    {"header": "底稿名称", "width": 30},
                    {"header": "所属循环", "width": 14},
                    {"header": "编制人", "width": 10},
                    {"header": "复核人", "width": 10},
                    {"header": "页数", "width": 8},
                    {"header": "备注", "width": 26},
                ],
                "rows": [["", "", "", "", "", None, ""] for _ in range(12)],
            }
        ],
    },
    "抽样表": {
        "title": "审计抽样表",
        "note": "记录总体规模、抽样方法、样本量与理由；样本量下调须注明原因并签字。",
        "sheets": [
            {
                "name": "抽样",
                "columns": [
                    {"header": "序号", "width": 7},
                    {"header": "凭证号 / 单据号", "width": 20},
                    {"header": "日期", "width": 13},
                    {"header": "摘要", "width": 28},
                    {"header": "金额", "width": 15, "fmt": "#,##0.00", "sum": True},
                    {"header": "科目", "width": 16},
                    {"header": "是否异常", "width": 10},
                    {"header": "说明", "width": 24},
                ],
                "rows": [[i, "", "", "", None, "", "", ""] for i in range(1, 21)],
                "total": True,
            }
        ],
    },
    "税负测算表": {
        "title": "税费测算表",
        "note": "税率与优惠必须带文号与执行期限（见税务线 references/01），差异写明理由。",
        "sheets": [
            {
                "name": "测算",
                "columns": [
                    {"header": "税种", "width": 18},
                    {"header": "计税依据", "width": 18, "fmt": "#,##0.00"},
                    {"header": "税率/征收率", "width": 14},
                    {"header": "应纳税额", "width": 16, "fmt": "#,##0.00", "sum": True},
                    {"header": "减免", "width": 14, "fmt": "#,##0.00"},
                    {"header": "实缴", "width": 14, "fmt": "#,##0.00"},
                    {"header": "差异", "width": 14, "fmt": "#,##0.00"},
                    {"header": "依据文号", "width": 26},
                ],
                "rows": [["增值税", None, "", None, None, None, None, ""],
                         ["城建税及附加", None, "", None, None, None, None, ""],
                         ["印花税", None, "", None, None, None, None, ""],
                         ["企业所得税", None, "", None, None, None, None, ""],
                         ["个人所得税", None, "", None, None, None, None, ""]],
                "total": True,
            }
        ],
    },
}


def main():
    ap = argparse.ArgumentParser(description="天信 Excel 表样生成器")
    ap.add_argument("--template", help=f"内置表样：{'、'.join(TEMPLATES)}")
    ap.add_argument("--spec", help="自定义 JSON 规格文件")
    ap.add_argument("--output", help="输出 .xlsx 路径")
    ap.add_argument("--verify", help="校验一个已生成的 .xlsx")
    args = ap.parse_args()

    if args.verify:
        wb = load_workbook(args.verify)
        total_formulas = 0
        for ws in wb.worksheets:
            formulas = sum(
                1
                for row in ws.iter_rows()
                for c in row
                if isinstance(c.value, str) and c.value.startswith("=")
            )
            total_formulas += formulas
            print(f"  {ws.title}: {ws.max_row} 行 × {ws.max_column} 列，公式 {formulas} 个，"
                  f"冻结 {ws.freeze_panes or '无'}")
        print(f"合计公式 {total_formulas} 个（Excel 打开后自动计算）")
        return

    if args.spec:
        with open(args.spec, encoding="utf-8") as f:
            spec = json.load(f)
    elif args.template:
        if args.template not in TEMPLATES:
            sys.exit(f"没有这个表样：{args.template}；可选：{'、'.join(TEMPLATES)}")
        spec = TEMPLATES[args.template]
    else:
        sys.exit("要么 --template，要么 --spec，二者必给其一")

    if not args.output:
        sys.exit("--output 必填")

    wb = Workbook()
    sheets = spec.get("sheets") or []
    for i, sheet in enumerate(sheets):
        # 包级 title / note 作为默认值下发到每张工作表（表内同名键优先）
        merged = dict(sheet)
        if spec.get("title"):
            merged.setdefault("title", spec["title"])
        if spec.get("note"):
            merged.setdefault("note", spec["note"])
        build_sheet(wb, merged, first=(i == 0))
    wb.save(args.output)
    sheets = "、".join(s.title for s in wb.worksheets)
    print(f"已生成 {args.output}（工作表：{sheets}）")


if __name__ == "__main__":
    main()
