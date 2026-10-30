#!/usr/bin/env python3
"""
Markdown → 天信 Word 文档（基于空壳模板）

与申万那套的差别：
  · 模板打开后清空正文、保留 sectPr（页眉 logo / 页脚页码域不重写）
  · 表格默认三线表（事务所报告的专业标志），不是全框线网格
  · 【事实】/【行业】/【推演】三级标记可选加底色，默认只留文字标记
  · 列表用手工前缀 + 悬挂缩进，不用 Word 自动编号（换机器不会错乱）

用法：
  python md2docx.py --input 报告.md --output 报告.docx \
      [--format report|gongwen] [--docno "天信咨字〔2026〕第 012 号"] \
      [--marks] [--title "文档主标题"]
"""

import argparse
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Mm

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from docx_helpers import (  # noqa: E402
    add_paragraph_border, repeat_table_header, set_cell_borders, set_rfonts,
    set_table_three_line, shade_cell, shade_run,
)

SKILL = os.path.dirname(HERE)
ASSETS = os.path.join(SKILL, 'assets')
TEMPLATES = {
    'report': os.path.join(ASSETS, 'tianxin-report.docx'),
    'gongwen': os.path.join(ASSETS, 'tianxin-gongwen.docx'),
}

NAVY = RGBColor(0x1D, 0x3A, 0x5E)
MUTED = RGBColor(0x52, 0x52, 0x52)
HEAD_FILL = 'F4F4F4'
MARK_FILL = {'事实': 'F4F4F4', '行业': 'EDF5FF', '推演': 'FDF3DA'}

LIST_RE = re.compile(r'^(\s*)([-*+]|(\d+)[.、)])\s+(.*)$')
TABLE_SEP_RE = re.compile(r'^\s*\|?[\s:|-]+\|?\s*$')


# ── 基础工具 ──────────────────────────────────────────────

def clear_body_keep_sectpr(doc):
    body = doc.element.body
    ns_w = '{' + body.nsmap['w'] + '}'
    sectpr = body.find(ns_w + 'sectPr')
    for child in list(body):
        body.remove(child)
    if sectpr is not None:
        body.append(sectpr)


def no_first_line_indent(para):
    """彻底取消首行缩进。

    只设 first_line_indent = Pt(0) 是不够的：样式里带着 w:firstLineChars="200"，
    而 OOXML 规定 firstLineChars 与 firstLine 同时存在时**前者优先**，
    于是段落级的 0 被样式里的 2 字符压回去——引用块、代码块、图片段落
    看起来都缩进了两个字符，改段落属性却怎么都不生效。
    """
    pf = para.paragraph_format
    pf.first_line_indent = Pt(0)
    ppr = para._p.get_or_add_pPr()
    ind = ppr.find(qn('w:ind'))
    if ind is not None:
        ind.set(qn('w:firstLineChars'), '0')


def hanging_indent(para, left_pt, hang_pt):
    """悬挂缩进，同样要把 firstLineChars 压成 0"""
    pf = para.paragraph_format
    pf.left_indent = Pt(left_pt)
    pf.first_line_indent = Pt(-hang_pt)
    ppr = para._p.get_or_add_pPr()
    ind = ppr.find(qn('w:ind'))
    if ind is not None:
        ind.set(qn('w:firstLineChars'), '0')


# ── 行内解析（**加粗** + 三级标记）────────────────────────

def add_inline(para, text, mark_shade=False, base_bold=False):
    # 表格单元格与引用块里常写 <br> 强制换行（写 SWOT、要列清单时几乎必用）。
    # 不处理的话它会以字面 "<br>" 印在纸上。Word 里对应的就是软换行 \n。
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.I)
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        bold = base_bold
        seg = part
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            seg = part[2:-2]
            bold = True
        # 标记切分
        if mark_shade and any(f'【{k}】' in seg for k in MARK_FILL):
            chunks = re.split(r'(【(?:事实|行业|推演)】)', seg)
            for ch in chunks:
                if not ch:
                    continue
                m = re.fullmatch(r'【(事实|行业|推演)】', ch)
                run = para.add_run(ch)
                run.bold = bold
                if m:
                    shade_run(run, MARK_FILL[m.group(1)])
            continue
        run = para.add_run(seg)
        run.bold = bold


# ── Markdown 解析 ────────────────────────────────────────

def parse_blocks(md):
    lines = md.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        raw = line.rstrip()

        if not raw.strip():
            i += 1
            continue

        # 代码块（含用制表符画的 ASCII 架构图）
        m = re.match(r'^\s*(`{3,}|~{3,})\s*(\S*)\s*$', raw)
        if m:
            fence, lang = m.group(1)[0], m.group(2)
            body, i = [], i + 1
            while i < len(lines) and not re.fullmatch(
                    r'\s*' + re.escape(fence) + r'{3,}\s*', lines[i]):
                body.append(lines[i])
                i += 1
            i += 1                      # 跳过收尾围栏
            blocks.append(('code', (lang, '\n'.join(body))))
            continue

        # 强制分页：只能用显式标记。
        # `---` 曾经也当分页用，但真实报告里它几乎总是「章与章之间的分隔线」，
        # 一份 13 章的报告因此被切成 41 页、其中 21 页只有小半内容。
        # 分页是排版意图，必须写清楚，不能靠一条横线猜。
        if re.fullmatch(r'\s*(:::pagebreak|<!--\s*pagebreak\s*-->|\\pagebreak)\s*', raw):
            blocks.append(('pagebreak', None))
            i += 1
            continue

        # 水平分隔线：只是视觉分隔，不换页。
        # 紧跟在章标题前面的那一条直接丢掉——章本来就另起一页，
        # 在上一页末尾留一条灰线纯属噪音（真实报告里 15/16 条都属这种）。
        if re.fullmatch(r'\s*(-{3,}|\*{3,}|_{3,})\s*', raw):
            nxt = next((x for x in lines[i + 1:] if x.strip()), '')
            if not re.match(r'^#{1,2}\s', nxt):
                blocks.append(('hr', None))
            i += 1
            continue

        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', raw)
        if m:
            blocks.append((f'h{len(m.group(1))}', m.group(2).strip()))
            i += 1
            continue

        # 图片
        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', raw.strip())
        if m:
            blocks.append(('image', (m.group(1).strip(), m.group(2).strip())))
            i += 1
            continue

        # 表格
        if raw.lstrip().startswith('|') and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1]):
            header = [c.strip() for c in raw.strip().strip('|').split('|')]
            rows, j = [], i + 2
            while j < len(lines) and lines[j].lstrip().startswith('|'):
                rows.append([c.strip() for c in lines[j].strip().strip('|').split('|')])
                j += 1
            blocks.append(('table', (header, rows)))
            i = j
            continue

        # 引用
        if raw.lstrip().startswith('>'):
            buf = []
            while i < len(lines) and lines[i].lstrip().startswith('>'):
                buf.append(re.sub(r'^\s*>\s?', '', lines[i]))
                i += 1
            blocks.append(('quote', '\n'.join(buf).strip()))
            continue

        # 列表
        m = LIST_RE.match(raw)
        if m:
            indent = len(m.group(2)) and len(m.group(1)) // 2
            ordered = m.group(3) is not None
            items = []
            while i < len(lines):
                mm = LIST_RE.match(lines[i])
                if not mm:
                    break
                lvl = len(mm.group(1)) // 2
                items.append((lvl, mm.group(2), mm.group(4).strip()))
                i += 1
            blocks.append(('list', items))
            continue

        # 数据来源 / 表题 / 图注 特殊段
        if re.match(r'^\s*(数据来源|资料来源|来源)\s*[：:]', raw):
            blocks.append(('source', raw.strip()))
            i += 1
            continue
        if re.match(r'^\s*\*\*(表|图)\s*\d+', raw):
            blocks.append(('caption', raw.strip()))
            i += 1
            continue

        # 普通段落（合并连续行）
        buf = [raw.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r'^(#{1,4}\s|\s*\||\s*>|\s*[-*+]\s|\s*\d+[.、)]\s|'
                r'\s*(`{3,}|~{3,})|'
                r'\s*([-*_]{3,}|:::pagebreak|<!--\s*pagebreak\s*-->|\\pagebreak)\s*$|!\[)',
                lines[i]):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(('p', ' '.join(buf)))
    return blocks


# ── 渲染 ─────────────────────────────────────────────────

USABLE_PT = 442          # A4 版心宽度：210 − 28 − 26 = 156mm ≈ 442pt


def _fit_title(para, text, max_pt=24, min_pt=17):
    """按字数把主标题缩到一行放得下（中文按 1 字宽、西文按 0.55 字宽计）"""
    units = sum(0.55 if ord(c) < 0x2E80 else 1.0 for c in text)
    if units <= 0:
        return
    size = min(max_pt, USABLE_PT / units)
    if size < max_pt:
        size = max(min_pt, size)
        for run in para.runs:
            run.font.size = Pt(size)


def render(doc, blocks, mark_shade=False, page_break_style='page', chapter_break=False):
    seen_h2 = False
    for kind, payload in blocks:
        if kind == 'h1':
            p = doc.add_paragraph(payload, style='Title')
            # 主标题按字数收字号：24pt 在 156mm 版心里只放得下约 16 个汉字，
            # 再长就会把最后一个字挤成单独一行（"……尽调报/告"），很刺眼。
            _fit_title(p, payload)
        elif kind == 'h2':
            p = doc.add_paragraph(payload, style='Heading 1')
            # 章另起页。`---` 不再兼任分页之后，章的分页由标题自己决定——
            # 这是它本来的语义：一章开始，就是新的一页开始。
            # 第一个章不另起：它前面的文档抬头本来就该和正文同页。
            if chapter_break and seen_h2:
                p.paragraph_format.page_break_before = True
            seen_h2 = True
        elif kind == 'h3':
            doc.add_paragraph(payload, style='Heading 2')
        elif kind == 'h4':
            doc.add_paragraph(payload, style='Heading 3')

        elif kind == 'p':
            p = doc.add_paragraph(style='Normal')
            add_inline(p, payload, mark_shade)

        elif kind == 'quote':
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.left_indent = Pt(21)
            no_first_line_indent(p)
            add_paragraph_border(p, edge='left', color='E0E0E0', sz=12, space=8)
            add_inline(p, payload, mark_shade)

        elif kind == 'image':
            alt, path = payload
            p = doc.add_paragraph(style='Normal')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            no_first_line_indent(p)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            if os.path.exists(path):
                p.add_run().add_picture(path, width=Mm(140))
                if alt:
                    c = doc.add_paragraph(style='天信图注')
                    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_inline(c, alt, False)
            else:
                r = p.add_run(f'［缺图：{alt or path}］')
                r.font.color.rgb = RGBColor(0xDA, 0x1E, 0x28)

        elif kind == 'list':
            counters = {}
            for lvl, marker, text in payload:
                ordered = bool(re.match(r'\d+', marker))
                if ordered:
                    counters[lvl] = counters.get(lvl, 0) + 1
                    prefix = f'{counters[lvl]}. '
                else:
                    prefix = '· '
                p = doc.add_paragraph(style='Normal')
                hanging_indent(p, 21 * (lvl + 1), 10.5)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.add_run(prefix)
                add_inline(p, text, mark_shade)

        elif kind == 'caption':
            p = doc.add_paragraph(style='天信表题')
            add_inline(p, payload, False, base_bold=True)

        elif kind == 'source':
            p = doc.add_paragraph(style='天信来源')
            add_inline(p, payload, False)

        elif kind == 'table':
            header, rows = payload
            ncol = len(header)
            table = doc.add_table(rows=1, cols=ncol)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            hdr = table.rows[0]
            for k, txt in enumerate(header):
                cell = hdr.cells[k]
                cell.text = ''
                p = cell.paragraphs[0]
                p.style = doc.styles['天信表格文字']
                add_inline(p, txt, False, base_bold=True)
                shade_cell(cell, HEAD_FILL)
            for row in rows:
                cells = table.add_row().cells
                for k in range(ncol):
                    txt = row[k] if k < len(row) else ''
                    p = cells[k].paragraphs[0]
                    p.style = doc.styles['天信表格文字']
                    add_inline(p, txt, mark_shade)
            set_table_three_line(table)
            for cell in table.rows[0].cells:
                set_cell_borders(cell, bottom=('single', 4, '1D3A5E'))
            repeat_table_header(table)
            # 表后留白
            trailing = doc.add_paragraph(style='Normal')
            trailing.paragraph_format.space_after = Pt(0)
            trailing.paragraph_format.line_spacing = Pt(6)
            no_first_line_indent(trailing)

        elif kind == 'code':
            lang, body = payload
            # ASCII 架构图必须靠等宽字体才对得齐：西文 Consolas、中文宋体
            # （宋体汉字正好两个半角宽，与框线字符的宽度假设一致）。
            lines_ = body.split('\n') if body else ['']
            for n, ln in enumerate(lines_):
                p = doc.add_paragraph(style='天信代码')
                no_first_line_indent(p)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                if n == 0:
                    p.paragraph_format.space_before = Pt(6)
                if n == len(lines_) - 1:
                    p.paragraph_format.space_after = Pt(8)
                p.add_run(ln if ln else ' ')

        elif kind == 'hr':
            # 分隔线：只是一条细线，不换页。用空段落 + 下边框，高度可控。
            p = doc.add_paragraph(style='Normal')
            pf = p.paragraph_format
            pf.space_before = Pt(10)
            pf.space_after = Pt(10)
            pf.line_spacing = Pt(1)
            no_first_line_indent(p)
            add_paragraph_border(p, edge='bottom', color='D9D9D9', sz=4, space=1)

        elif kind == 'pagebreak':
            doc.add_page_break()


# ── 主流程 ───────────────────────────────────────────────

def convert(input_md, output_docx, fmt='report', docno=None, marks=False, title=None,
            keep_template_body=False, chapter_break=None):
    tpl = TEMPLATES[fmt]
    if not os.path.exists(tpl):
        sys.exit(f'✗ 找不到模板 {tpl}，先运行 build_template.py')

    with open(input_md, encoding='utf-8') as f:
        md = f.read()

    doc = Document(tpl)

    if not keep_template_body:
        clear_body_keep_sectpr(doc)

    # 换页眉报告编号
    if docno is not None:
        sec = doc.sections[0]
        for para in sec.header.paragraphs:
            for run in para.runs:
                if '天信' in run.text and '字' in run.text:
                    run.text = docno

    if title:
        p = doc.add_paragraph(title, style='Title')
        p.paragraph_format.first_line_indent = Pt(0)

    blocks = parse_blocks(md)
    if chapter_break is None:
        n_h2 = sum(1 for k, _ in blocks if k == 'h2')
        chapter_break = n_h2 >= 6 and len(md) >= 8000
    render(doc, blocks, mark_shade=marks, chapter_break=chapter_break)

    # 去掉尾部空段
    body = doc.element.body
    for child in list(body)[::-1]:
        if child.tag.endswith('}p') and not ''.join(child.itertext()).strip():
            body.remove(child)
        else:
            break

    doc.save(output_docx)
    n_par = len(doc.paragraphs)
    n_tbl = len(doc.tables)
    print(f'✅ {output_docx}')
    print(f'   段落 {n_par} · 表格 {n_tbl} · 标记底色 {"开" if marks else "关"} · 版式 {fmt}')
    return output_docx


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True)
    ap.add_argument('--output', required=True)
    ap.add_argument('--format', default='report', choices=list(TEMPLATES))
    ap.add_argument('--docno', default=None)
    ap.add_argument('--title', default=None)
    ap.add_argument('--marks', action='store_true')
    g = ap.add_mutually_exclusive_group()
    g.add_argument('--chapter-break', dest='cb', action='store_true', default=None,
                   help='每个 h2（章）另起一页；默认按体量自动判断')
    g.add_argument('--no-chapter-break', dest='cb', action='store_false',
                   help='连续排版，不按章分页')
    a = ap.parse_args()
    convert(a.input, a.output, a.format, a.docno, a.marks, a.title,
            chapter_break=a.cb)
